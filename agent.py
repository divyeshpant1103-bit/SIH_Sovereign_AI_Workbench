"""
The agent. Two model calls, everything numeric done deterministically.

Usage:
    python agent.py data/inputs/INS-2026-0847_P-101-A_deviation.pdf

Flow:
    1. read the report PDF                              [no model]
    2. model call 1: extract facts as JSON              [model - reading]
    3. evaluate measurements against criteria           [no model - arithmetic]
    4. retrieve evidence for the observations           [no model]
    5. model call 2: interpret observations, recommend  [model - judgment]
    6. render .docx                                     [no model]
    7. reopen and validate                              [no model]

Pass/fail against a stated numeric limit is arithmetic, not judgment, so it is
never model-generated. A hallucinated pass/fail in a refinery is the most
dangerous output this system could produce.
"""

import sys
import time
from pathlib import Path

import pymupdf
import requests
from docx import Document
from docx.shared import Pt, RGBColor

from criteria import evaluate, summarise
from schemas import Assessment, ReportFacts
from search import Retriever
from trace import Trace

OLLAMA_URL = "http://localhost:11434/api/chat"
MODEL = "workbench"
CODE_MODEL = "qwen2.5-coder:1.5b"
OUTPUT_DIR = Path("data/outputs")

MAX_SECONDS = 180
CALL_TIMEOUT = 120
MIN_TEXT_CHARS = 50  # below this, treat the PDF as having no usable text layer


class CorruptPDFError(Exception):
    """Raised when the PDF cannot be opened at all."""


class NoTextLayerError(Exception):
    """Raised when the PDF opens but yields no extractable text (scanned page)."""


# ----------------------------------------------------------------------
# model access
# ----------------------------------------------------------------------

def call_model(system, user, schema, model=MODEL, trace=None, label=""):
    """
    One constrained call. The 'format' parameter forces output to match the
    JSON schema at decode time, so the model cannot return invalid JSON.
    """
    body = {
        "model": model,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        "format": schema.model_json_schema(),
        "stream": False,
        "options": {"temperature": 0.1, "num_ctx": 8192},
    }

    started = time.time()
    resp = requests.post(OLLAMA_URL, json=body, timeout=CALL_TIMEOUT)
    resp.raise_for_status()
    content = resp.json()["message"]["content"]
    elapsed = round(time.time() - started, 1)

    if trace:
        trace.emit("model_call", label=label, model=model, seconds=elapsed)

    return schema.model_validate_json(content)


def select_model(request_text, trace):
    """Deterministic router. A lookup rule, not an AI choosing an AI."""
    lowered = request_text.lower()
    if any(w in lowered for w in ("calculate", "compute", "script", "code", "formula")):
        choice, reason = CODE_MODEL, "request contains calculation or code intent"
    else:
        choice, reason = MODEL, "document comparison task, general model"
    trace.emit("model_selected", model=choice, reason=reason)
    return choice


# ----------------------------------------------------------------------
# steps
# ----------------------------------------------------------------------

def read_pdf(path, trace):
    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise CorruptPDFError(str(exc)) from exc

    text = "\n".join(page.get_text() for page in doc)
    pages = doc.page_count
    doc.close()

    if len(text.strip()) < MIN_TEXT_CHARS:
        raise NoTextLayerError("no extractable text layer")

    trace.emit("document_read", file=Path(path).name, pages=pages, chars=len(text))
    return text


EXTRACT_SYSTEM = """You read industrial equipment inspection reports and extract the
facts exactly as written. Do not interpret, judge, or add anything. Copy values
verbatim including units and any fluctuation figures. If a field is absent, write
"not stated"."""


def extract_facts(report_text, model, trace):
    facts = call_model(
        EXTRACT_SYSTEM,
        f"Extract the facts from this inspection report.\n\n{report_text}",
        ReportFacts, model=model, trace=trace, label="extract_facts",
    )
    trace.emit("facts_extracted", equipment=facts.equipment_tag,
               measurements=len(facts.measurements),
               observations=len(facts.observations))
    return facts


def gather_evidence(facts, retriever, trace):
    """One search, built from the report's own observations, so symptom
    descriptions can reach the technical handbook."""
    if not facts.observations:
        return []

    query = " ".join(facts.observations)[:300]
    result = retriever.search(query)
    trace.emit("retrieval", query=query[:70], refused=result.refused,
               top_score=round(result.top_score, 3), hits=len(result.evidence))

    if result.refused:
        return []

    citations = sorted({e.cite() for e in result.evidence})
    trace.emit("evidence_gathered", count=len(result.evidence),
               sources=", ".join(citations)[:150], citations=citations)
    return result.evidence


ASSESS_SYSTEM = """You are a rotating equipment engineer reviewing an inspection.

The numeric comparison against acceptance limits has ALREADY been done and is given
to you. Do not repeat it, do not contradict it, and do not list measured values again.

Your job is narrower:

1. pattern_identified: If the inspector's OBSERVATIONS match a recognised failure
   pattern described in the technical evidence, write one or two sentences naming it
   and explaining which observations support it. Otherwise write "none identified".

2. advisory_findings: ONLY items from the observations that have no numeric limit and
   are NOT already in the deviations list. Typically maintenance issues such as missing
   fasteners or overdue logs. Usually one or two. If there are none, return an empty list.

3. recommendation: What should happen next, in two to four sentences.

Citation rules, follow exactly:
- source_doc must be a document name copied from a SOURCE line, nothing else.
- source_page must be a short page label copied from a SOURCE line, such as
  "ME-03 p.12" or "p.2". Never put a sentence in source_page.
- Use ONLY the evidence provided. Never use outside knowledge."""


def assess(facts, findings, counts, evidence, model, trace):
    ev_block = "\n\n".join(
        f"[{i}] SOURCE: {e.doc}, {e.page_label}\n{e.text[:900]}"
        for i, e in enumerate(evidence, 1)
    ) or "No supporting technical evidence retrieved."

    breaches = "\n".join(
        f"- {f['parameter']}: {f['measured_value']} [{f['severity'].upper()}]"
        for f in findings if f["severity"] in ("action", "alert")
    ) or "- none"

    observations = "\n".join(f"- {o}" for o in facts.observations)

    user = f"""EQUIPMENT: {facts.equipment_tag} ({facts.equipment_description})
REPORT: {facts.report_number}, {facts.inspection_date}

ALREADY DETERMINED - deviations against acceptance limits:
{breaches}

INSPECTOR OBSERVATIONS
{observations}

TECHNICAL EVIDENCE
{ev_block}

Identify any failure pattern the observations match, and recommend next steps."""

    result = call_model(ASSESS_SYSTEM, user, Assessment,
                        model=model, trace=trace, label="assess")
    trace.emit("assessment_complete", pattern=result.pattern_identified[:60],
               advisories=len(result.advisory_findings))
    return result


# ----------------------------------------------------------------------
# deliverable
# ----------------------------------------------------------------------

SEV_COLOR = {
    "action": RGBColor(0xA8, 0x50, 0x1E),
    "alert": RGBColor(0xB8, 0x86, 0x0B),
    "acceptable": RGBColor(0x2E, 0x6B, 0x3E),
}


def build_docx(facts, findings, status, counts, assessment, unmatched, out_path, trace):
    """Deterministic. The model supplied structured data; this code builds the file."""
    doc = Document()
    doc.add_heading("Equipment Approval Note", level=0)

    p = doc.add_paragraph()
    r = p.add_run("Generated by KAVACH sovereign workbench. Numeric comparisons are "
                  "deterministic. Requires human review and approval before use.")
    r.italic = True
    r.font.color.rgb = RGBColor(0xA8, 0x50, 0x1E)

    doc.add_heading("Equipment", level=1)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in [
        ("Equipment tag", facts.equipment_tag),
        ("Description", facts.equipment_description),
        ("Report number", facts.report_number),
        ("Inspection date", facts.inspection_date),
        ("Overall status", status.replace("_", " ")),
        ("Summary", f"{counts['action']} action, {counts['alert']} alert, "
                    f"{counts['acceptable']} acceptable"),
    ]:
        row = t.add_row().cells
        row[0].text = k
        row[1].paragraphs[0].add_run(str(v)).bold = True

    doc.add_heading("Findings against acceptance criteria", level=1)
    for i, f in enumerate(findings, 1):
        h = doc.add_paragraph()
        run = h.add_run(f"{i}. {f['parameter']}  [{f['severity'].upper()}]")
        run.bold = True
        run.font.color.rgb = SEV_COLOR.get(f["severity"], RGBColor(0, 0, 0))

        doc.add_paragraph(f"Measured: {f['measured_value']}", style="List Bullet")
        doc.add_paragraph(f"Criterion: {f['applicable_limit']}", style="List Bullet")
        doc.add_paragraph(f["description"])

        src = doc.add_paragraph()
        s = src.add_run(f"Source: {f['source_doc']}, {f['source_page']}")
        s.italic = True
        s.font.size = Pt(9)

    if assessment.pattern_identified and \
            assessment.pattern_identified.lower() not in ("none", "none identified"):
        doc.add_heading("Engineering assessment", level=1)
        text = assessment.pattern_identified.strip()
        if len(text.split()) < 4:
            text = (f"Pattern identified: {text}. The inspector's observations match "
                    f"this pattern as described in the retrieved technical evidence.")
        doc.add_paragraph(text)

    if assessment.advisory_findings:
        doc.add_heading("Advisory", level=1)
        for a in assessment.advisory_findings:
            doc.add_paragraph(a.description, style="List Bullet")
            page = a.source_page.strip()
            if len(page) > 30:
                page = page[:30].rsplit(" ", 1)[0] + "..."
            doc_name = a.source_doc.strip()
            if page and doc_name.endswith(page):
                doc_name = doc_name[: -len(page)].rstrip(" ,")
            src = doc.add_paragraph()
            s = src.add_run(f"Source: {doc_name}, {page}")
            s.italic = True
            s.font.size = Pt(9)

    if unmatched:
        doc.add_heading("Not covered by available criteria", level=1)
        doc.add_paragraph("No acceptance criterion was found in the knowledge base "
                          "for the following. These are recorded but not assessed.")
        for u in unmatched:
            doc.add_paragraph(u, style="List Bullet")

    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(assessment.recommendation)

    doc.add_heading("Provenance", level=1)
    doc.add_paragraph(f"Trace ID: {trace.id}")
    doc.add_paragraph("Numeric comparisons performed deterministically against "
                      "SOP-PMP-114 Section 2, not generated by a language model.")
    doc.add_paragraph("Generated entirely on local hardware. No external network "
                      "calls were made during this task.")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    trace.emit("artifact_created", file=out_path.name)
    return out_path


def validate_docx(path, trace):
    try:
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        checks = {
            "opens": True,
            "has_heading": "Approval Note" in text,
            "has_findings": "Findings against acceptance criteria" in text,
            "has_recommendation": "Recommendation" in text,
            "has_citations": "Source:" in text,
            "has_provenance": "Trace ID" in text,
        }
    except Exception as exc:
        trace.emit("artifact_invalid", error=str(exc)[:100])
        return False

    ok = all(checks.values())
    trace.emit("artifact_validated", valid=ok,
               failed=",".join(k for k, v in checks.items() if not v) or "none")
    return ok


# ----------------------------------------------------------------------
# orchestration
# ----------------------------------------------------------------------

def run(pdf_path, request="Check this inspection report against our procedures "
                          "and prepare an approval note.",
        on_event=None, retriever=None):
    """
    Runs the full flow and returns a plain dict describing the outcome, so
    callers (the CLI, the UI) can render it however they need to:

        {
            "trace_id": str,
            "refused": bool,          # True if no acceptance criteria applied
            "reason": str | None,     # why, when refused or aborted
            "status": str | None,     # "deviations_found" / "alerts_only" / "no_deviations"
            "counts": dict | None,    # {"action": n, "alert": n, "acceptable": n}
            "findings": list[dict],
            "unmatched": list[str],
            "assessment": Assessment | None,
            "out": Path | None,       # the generated .docx, if any
            "valid": bool | None,
            "elapsed": float,
        }

    `on_event` is forwarded to Trace() so a caller can stream events live.
    `retriever` lets a caller pass in an already-loaded Retriever (the
    embedding model is slow to load) instead of paying that cost every run.
    """
    trace = Trace(on_event=on_event)
    started = time.time()

    print(f"\nTrace {trace.id}")
    print("=" * 74)

    trace.emit("task_received", file=Path(pdf_path).name, request=request[:80])
    model = select_model(request, trace)
    report_text = read_pdf(pdf_path, trace)
    facts = extract_facts(report_text, model, trace)

    findings, unmatched = evaluate(facts.measurements)
    status, counts = summarise(findings)
    trace.emit("criteria_evaluated", status=status, **counts,
               unmatched=len(unmatched))

    if not findings:
        reason = ("no applicable acceptance criteria in the knowledge base "
                  "for this equipment type")
        trace.emit("refused", reason=reason)
        print("\nREFUSED - no applicable acceptance criteria for this equipment.")
        print(f"Elapsed {round(time.time() - started, 1)}s")
        return {
            "trace_id": trace.id, "refused": True, "reason": reason,
            "status": None, "counts": None, "findings": [], "unmatched": unmatched,
            "assessment": None, "out": None, "valid": None,
            "elapsed": round(time.time() - started, 1),
        }

    if retriever is None:
        retriever = Retriever()
    evidence = gather_evidence(facts, retriever, trace)

    if time.time() - started > MAX_SECONDS:
        reason = "time budget exceeded"
        trace.emit("aborted", reason=reason)
        return {
            "trace_id": trace.id, "refused": False, "reason": reason,
            "status": status, "counts": counts, "findings": findings,
            "unmatched": unmatched, "assessment": None, "out": None, "valid": None,
            "elapsed": round(time.time() - started, 1),
        }

    assessment = assess(facts, findings, counts, evidence, model, trace)

    out = OUTPUT_DIR / f"ApprovalNote_{facts.equipment_tag}_{trace.id}.docx"
    build_docx(facts, findings, status, counts, assessment, unmatched, out, trace)
    valid = validate_docx(out, trace)

    total = round(time.time() - started, 1)
    trace.emit("task_complete", seconds=total, artifact=out.name, valid=valid)

    print("=" * 74)
    print(f"Status:   {status}   ({counts['action']} action, {counts['alert']} alert, "
          f"{counts['acceptable']} acceptable)")
    for f in findings:
        print(f"  [{f['severity']:>10}] {f['parameter'][:44]:<44} {f['measured_value'][:26]}")
    if unmatched:
        print(f"  not assessed: {', '.join(unmatched)}")
    print(f"Pattern:  {assessment.pattern_identified[:100]}")
    print(f"Artifact: {out}  (valid={valid})")
    print(f"Elapsed:  {total}s")

    return {
        "trace_id": trace.id, "refused": False, "reason": None,
        "status": status, "counts": counts, "findings": findings,
        "unmatched": unmatched, "assessment": assessment, "out": out,
        "valid": valid, "elapsed": total,
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python agent.py <path-to-inspection-report.pdf>")
        sys.exit(1)
    run(sys.argv[1])